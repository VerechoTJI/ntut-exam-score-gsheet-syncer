"""Read student identities, never grade columns, from an Excel roster."""
import math
import re

from openpyxl import load_workbook

from exam_config import require

HEADERS = ('班級', '姓名', '學號')


def header_values(worksheet):
    return [str(cell.value or '').strip() for cell in next(worksheet.iter_rows(), ())]


def student_id(cell):
    value = cell.value
    if type(value) in (int, float):
        require(math.isfinite(value) and value == int(value), 'Student ID must be an integer or text')
        text = str(int(value))
        if re.fullmatch('0+', cell.number_format):
            text = text.zfill(len(cell.number_format))
        return text
    require(isinstance(value, str) and value.strip() and not value.startswith('='), 'Missing or invalid student ID')
    return value.strip()


def read_roster(path, sheet=None):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet:
            require(sheet in workbook.sheetnames, f'Roster worksheet {sheet!r} not found')
            worksheet = workbook[sheet]
        else:
            matches = [ws for ws in workbook.worksheets if all(h in header_values(ws) for h in HEADERS)]
            require(len(matches) == 1, 'Missing or ambiguous roster headers; specify --roster-sheet')
            worksheet = matches[0]
        headers = header_values(worksheet)
        require(all(headers.count(h) == 1 for h in HEADERS), 'Required roster headers must each occur once')
        columns = [headers.index(h) for h in HEADERS]
        identities = []
        seen = set()
        for row in worksheet.iter_rows(min_row=2, max_col=len(headers)):
            selected = [row[index] for index in columns]
            if all(cell.value in (None, '') for cell in selected):
                continue
            klass, name = [str(cell.value or '').strip() for cell in selected[:2]]
            require(klass and name and not name.startswith('='), 'Roster class and name required')
            sid = student_id(selected[2])
            require(sid not in seen, f'Duplicate student ID at Excel row {selected[2].row}')
            seen.add(sid)
            identities.append({'student_id': sid, 'name': name, 'class': klass})
        require(identities, 'Roster contains no students')
        return identities
    finally:
        workbook.close()
